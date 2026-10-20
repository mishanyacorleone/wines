"""Выгрузка аналитики в Excel: тот же отчёт, что на дашборде, плюс сами события периода.

Листы: «Сводка» (показатели), «По дням» (сканы и график), «Воронка», «Вина»,
«Регионы и производители», «Поделиться», «Вовлечение», «Точность распознавания», «Оценки», «Сети», «События». Числа — числами,
доли — процентным форматом, итоги — формулами: файл можно дальше считать в Excel.
"""

from __future__ import annotations

import io
import json
from datetime import date, datetime, timezone

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from . import report
from .db import Event

# Сырых событий в файле — не больше этого: дальше Excel неудобен, а стенду хватает с запасом
MAX_EVENTS = 50_000

WINE = "7A1F3D"
HEAD_FILL = PatternFill("solid", fgColor=WINE)
HEAD_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color=WINE)
PERCENT = "0.0%"

# Показатели «Сводки»: ключ kpi → (подпись, формат)
KPI_ROWS = [
    ("scans", "Сканов", "#,##0"),
    ("found", "Вино найдено", "#,##0"),
    ("not_found", "Вина нет в каталоге или не распознано", "#,##0"),
    ("not_found_share", "Доля «не найдено»", PERCENT),
    ("visitors", "Посетителей (устройств)", "#,##0"),
    ("visitors_with_account", "Из них с аккаунтом", "#,##0"),
    ("registrations", "Регистраций", "#,##0"),
    ("conversion", "Конверсия: сканировал без аккаунта → зарегистрировался", PERCENT),
    ("returning", "Вернулись хотя бы раз (активны 2+ дня)", "#,##0"),
    ("returning_share", "Доля вернувшихся", PERCENT),
    ("shares", "Поделились", "#,##0"),
    ("scan_p95_ms", "Время скана, p95, мс", "#,##0"),
]

EVENT_NAMES = {
    "page_view": "открыл страницу", "scan": "скан", "register": "регистрация", "share": "поделился",
    "nudge_shown": "предложение войти", "nudge_click": "нажал «войти»", "map_open": "открыл карту",
    "region_open": "открыл регион", "achievement": "достижение", "login": "вход",
    "quiz": "вопрос дня", "diary": "запись в дневник", "diary_prompt": "дневник: предложили войти",
    "gift_open": "открыл открытку", "gift_save": "сохранил подарок", "prize": "получил приз", "year_open": "итоги года",
    "auth_prompt": "предложили войти",
    "recognition": "ответ «Это ваше вино?»", "rate": "👍/👎 в истории", "taste": "«Мой вкус»",
    "winery": "винодельня",
}
# Как на дашборде (web/app/analytics.js, NAMES)
SHARE_NAMES = {"native": "Меню «Поделиться»", "tg": "Telegram", "vk": "ВКонтакте", "save": "Сохранили картинку",
               "wine": "Вино", "achievement": "Достижение", "map": "Винная карта", "gift": "Открытка"}


def build_xlsx(session: Session, days: int, include_demo: bool) -> bytes:
    data = report.build(session, days, include_demo=include_demo)
    wb = Workbook()
    _summary(wb.active, data, include_demo)
    _daily(wb.create_sheet("По дням"), data["daily"])
    _funnel(wb.create_sheet("Воронка"), data["funnel"])
    _table(wb.create_sheet("Вина"), "Самые сканируемые вина (найдено)", ["Вино", "Slug", "Сканов"],
           [[w["title"], w["slug"], w["count"]] for w in data["top_wines"]])
    places = wb.create_sheet("Регионы и производители")
    _table(places, "Регионы найденных вин", ["Регион", "Сканов"],
           [[r["name"], r["count"]] for r in data["top_regions"]])
    _table(places, "Производители", ["Производитель", "Сканов"],
           [[m["name"], m["count"]] for m in data["top_manufacturers"]], first_col=4)
    _table(places, "Винная карта: открывали регион", ["Регион", "Открытий"],
           [[r["name"], r["count"]] for r in data["map"]["regions"]], first_col=7,
           note=f"Карту открывали: {data['map']['opens']}")
    shares = wb.create_sheet("Поделиться")
    _table(shares, "Куда делились", ["Канал", "Раз"],
           [[SHARE_NAMES.get(c["name"], c["name"]), c["count"]] for c in data["share_channels"]])
    _table(shares, "Чем делились", ["Что", "Раз"],
           [[SHARE_NAMES.get(k["name"], k["name"]), k["count"]] for k in data["share_kinds"]], first_col=4)
    _engagement(wb.create_sheet("Вовлечение"), data["engagement"])
    _accuracy(wb.create_sheet("Точность распознавания"), data["feedback"])
    _ratings(wb.create_sheet("Оценки"), data["feedback"])
    _table(wb.create_sheet("Сети"), "Откуда сканируют (адрес без последнего октета)",
           ["Сеть", "Сканов", "Устройств"], [[i["network"], i["scans"], i["visitors"]] for i in data["ips"]],
           note=f"Доля сканов из локальной сети: {data['local_share']:.1%}")
    _events(wb.create_sheet("События"), session, data["period"]["from"], include_demo)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _engagement(ws: Worksheet, g: dict) -> None:
    """Вовлечение: показатели и «откуда регистрации» с конверсией по каждому призыву войти."""
    q, d, gf = g["quiz"], g["diary"], g["gifts"]
    _header(ws, 1, ["Показатель", "Значение"])
    rows = [
        ("Ответов на вопрос дня", q["answers"], "#,##0"),
        ("Игроков (устройств)", q["players"], "#,##0"),
        ("Доля верных ответов", q["correct_share"], PERCENT),
        ("Записей в дневнике «попробовал»", d["tasted"], "#,##0"),
        ("Записей «хочу попробовать»", d["want"], "#,##0"),
        ("Открыток создано", gf["created"], "#,##0"),
        ("Открыток открыто получателями", gf["opened"], "#,##0"),
        ("Сохранили подарок себе", gf["saved"], "#,##0"),
        ("Призов (промокодов) выдано", g["prizes"], "#,##0"),
        ("Открывали «Мой винный год»", g["year_opens"], "#,##0"),
        ("«Мой вкус»: открывали", g["taste"]["opens"], "#,##0"),
        ("«Мой вкус»: сохранили", g["taste"]["saves"], "#,##0"),
        ("Подсказок «понравится ли» показано", g["taste"]["matches"], "#,##0"),
        ("Винодельни: открывали страницу", g["winery"]["open"], "#,##0"),
        ("Винодельни: переходов к записи на тур", g["winery"]["tour"], "#,##0"),
        ("Винодельни: «Хочу съездить»", g["winery"]["want"], "#,##0"),
    ]
    for row, (label, value, fmt) in enumerate(rows, start=2):
        ws.cell(row, 1, label)
        ws.cell(row, 2, value).number_format = fmt
    top = len(rows) + 3
    ws.cell(top, 1, "Откуда регистрации — последний призыв войти перед регистрацией").font = Font(bold=True, color=WINE)
    _header(ws, top + 1, ["Призыв", "Увидели (устройств)", "Зарегистрировались", "Конверсия"])
    for row, s in enumerate(g["registration_sources"], start=top + 2):
        ws.cell(row, 1, s["name"])
        ws.cell(row, 2, s["prompted"] or None)
        ws.cell(row, 3, s["registrations"])
        ws.cell(row, 4, f"=IF(B{row}>0,C{row}/B{row},\"\")").number_format = PERCENT
    if not g["registration_sources"]:
        ws.cell(top + 2, 1, "нет данных за период").font = Font(italic=True, color="888888")
    _widths(ws, [52, 20, 20, 12])


def _accuracy(ws: Worksheet, f: dict) -> None:
    """Точность по ответам пользователей «Это ваше вино?» — по неделям, чтобы видеть её на длинном промежутке."""
    ws["A1"] = "Точность распознавания по ответам пользователей «Это ваше вино?»"
    ws["A1"].font = Font(bold=True, color=WINE)
    ws["A2"] = ("Точность = «Да, это оно» / все ответы на «нашли». Отвечает не каждый: это оценка на тех, кто ответил, "
                "а не на всех сканах.")
    ws["A2"].font = Font(italic=True, color="888888")
    _header(ws, 4, ["Неделя (с понедельника)", "Сканов", "Сканер: «нашли» — ответили", "Да, это оно",
                    "Нет — выбрали другое из похожих", "Нет — «моего вина нет»", "Точность",
                    "Сканер: «нет в каталоге» — выбрали из аналогов (было в каталоге)",
                    "Сканер: «нет в каталоге» — «моего вина нет» (верно)"])
    first = 5
    for row, w in enumerate(f["weeks"], start=first):
        m, n = w["match"], w["not_found"]
        ws.cell(row, 1, date.fromisoformat(w["week"])).number_format = "dd.mm.yyyy"
        ws.cell(row, 2, w["scans"])
        ws.cell(row, 3, f"=D{row}+E{row}+F{row}")
        ws.cell(row, 4, m["confirmed"])
        ws.cell(row, 5, m["picked"])
        ws.cell(row, 6, m["missing"])
        ws.cell(row, 7, f"=IF(C{row}>0,D{row}/C{row},\"\")").number_format = PERCENT
        ws.cell(row, 8, n["picked"])
        ws.cell(row, 9, n["missing"])
    last = first + len(f["weeks"]) - 1
    total = last + 1
    ws.cell(total, 1, "Итого").font = Font(bold=True)
    for col in "BCDEFHI":
        ws[f"{col}{total}"] = f"=SUM({col}{first}:{col}{last})" if f["weeks"] else 0
    ws[f"G{total}"] = f"=IF(C{total}>0,D{total}/C{total},\"\")"
    ws[f"G{total}"].number_format = PERCENT
    for col in "ABCDEFGHI":
        ws[f"{col}{total}"].font = Font(bold=True)
    if not f["weeks"]:
        ws.cell(first, 1, "нет данных за период").font = Font(italic=True, color="888888")
    _widths(ws, [16, 10, 16, 12, 18, 16, 11, 24, 22])
    ws.row_dimensions[4].height = 48
    ws.freeze_panes = "B5"
    if len(f["weeks"]) > 1:
        chart = LineChart()
        chart.title, chart.y_axis.title = "Точность по неделям", "Точность"
        chart.y_axis.number_format, chart.y_axis.scaling.min, chart.y_axis.scaling.max = "0%", 0, 1
        chart.add_data(Reference(ws, min_col=7, min_row=4, max_row=last), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
        chart.series[0].graphicalProperties.line.solidFill = WINE
        chart.height, chart.width, chart.legend = 8, 20, None
        ws.add_chart(chart, f"A{total + 3}")


def _ratings(ws: Worksheet, f: dict) -> None:
    """Вкусовые оценки: звёзды дневника и 👍/👎 в истории сканов (это «нравится», а не точность распознавания)."""
    ws["A1"] = "Оценки вин в дневнике"
    ws["A1"].font = Font(bold=True, color=WINE)
    _header(ws, 2, ["Оценка", "Записей", "Доля"])
    for row, s in enumerate(f["stars"], start=3):
        ws.cell(row, 1, "★" * s["stars"])
        ws.cell(row, 2, s["count"])
        ws.cell(row, 3, f"=IF(SUM($B$3:$B$7)>0,B{row}/SUM($B$3:$B$7),\"\")").number_format = PERCENT
    ws.cell(8, 1, "Средняя").font = Font(bold=True)
    ws["B8"] = "=IF(SUM(B3:B7)>0,(5*B3+4*B4+3*B5+2*B6+B7)/SUM(B3:B7),\"\")"
    ws["B8"].number_format, ws["B8"].font = "0.00", Font(bold=True)
    ws["A10"] = "👍/👎 в истории сканов (нажатия за период)"
    ws["A10"].font = Font(bold=True, color=WINE)
    _header(ws, 11, ["Оценка", "Нажатий"])
    ws["A12"], ws["B12"] = "👍 Нравится", f["likes"]
    ws["A13"], ws["B13"] = "👎 Не нравится", f["dislikes"]
    _widths(ws, [42, 12, 10])


def _summary(ws: Worksheet, data: dict, include_demo: bool) -> None:
    ws.title = "Сводка"
    period = data["period"]
    ws["A1"] = "Аналитика сканера «Своё Вино»"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Период: {period['from']} — {period['to']} ({period['days']} дн.)"
    ws["A3"] = ("В отчёте есть демо-данные (сгенерированы для показа), настоящие события — вместе с ними"
                if data["has_demo"] else
                "Только настоящие события" if not include_demo else "Демо-данных в периоде нет")
    if data["has_demo"]:
        ws["A3"].font = Font(italic=True, color="8A2E24")
    _header(ws, 5, ["Показатель", "Значение"])
    for row, (key, label, fmt) in enumerate(KPI_ROWS, start=6):
        ws.cell(row, 1, label)
        cell = ws.cell(row, 2, data["kpi"].get(key))
        cell.number_format = fmt
    _widths(ws, [58, 16])


def _daily(ws: Worksheet, daily: list[dict]) -> None:
    _header(ws, 1, ["Дата", "Найдено", "Не найдено", "Всего сканов"])
    for row, d in enumerate(daily, start=2):
        ws.cell(row, 1, date.fromisoformat(d["date"])).number_format = "dd.mm.yyyy"
        ws.cell(row, 2, d["found"])
        ws.cell(row, 3, d["not_found"])
        ws.cell(row, 4, f"=B{row}+C{row}")
    last = len(daily) + 1
    total = last + 1
    ws.cell(total, 1, "Итого").font = Font(bold=True)
    for col in "BCD":
        ws[f"{col}{total}"] = f"=SUM({col}2:{col}{last})"
        ws[f"{col}{total}"].font = Font(bold=True)
    _widths(ws, [14, 12, 14, 14])
    ws.freeze_panes = "A2"
    if daily:
        chart = BarChart()
        chart.type, chart.grouping, chart.overlap = "col", "stacked", 100
        chart.title, chart.y_axis.title = "Сканы по дням", "Сканов"
        chart.add_data(Reference(ws, min_col=2, max_col=3, min_row=1, max_row=last), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=last))
        for series, color in zip(chart.series, (WINE, "D7B9C4")):  # найдено — цвет сайта, не найдено — светлее
            series.graphicalProperties.solidFill = color
            series.graphicalProperties.line.solidFill = color
        chart.height, chart.width = 9, 22
        ws.add_chart(chart, "F2")


def _funnel(ws: Worksheet, funnel: list[dict]) -> None:
    _header(ws, 1, ["Шаг", "Устройств", "От первого шага"])
    for row, step in enumerate(funnel, start=2):
        ws.cell(row, 1, step["step"])
        ws.cell(row, 2, step["value"])
        cell = ws.cell(row, 3, f"=IF($B$2=0,0,B{row}/$B$2)")
        cell.number_format = PERCENT
    _widths(ws, [46, 12, 16])


def _table(ws: Worksheet, title: str, head: list[str], rows: list[list], *, first_col: int = 1,
           note: str | None = None) -> None:
    ws.cell(1, first_col, title).font = Font(bold=True, color=WINE)
    _header(ws, 2, head, first_col)
    for r, values in enumerate(rows, start=3):
        for c, value in enumerate(values, start=first_col):
            ws.cell(r, c, value)
    if not rows:
        ws.cell(3, first_col, "нет данных за период").font = Font(italic=True, color="888888")
    if note:
        ws.cell(max(len(rows), 1) + 4, first_col, note).font = Font(italic=True)
    for offset, width in enumerate([34] + [12] * (len(head) - 1)):
        ws.column_dimensions[get_column_letter(first_col + offset)].width = width


def _events(ws: Worksheet, session: Session, start: str, include_demo: bool) -> None:
    head = ["Время (UTC)", "Событие", "Устройство", "Пользователь", "Сеть", "Демо", "Подробности"]
    _header(ws, 1, head)
    since = datetime.combine(date.fromisoformat(start), datetime.min.time(), tzinfo=timezone.utc)
    query = select(Event).where(Event.ts >= since).order_by(Event.ts).limit(MAX_EVENTS + 1)
    if not include_demo:
        query = query.where(Event.demo.is_(False))
    events = session.scalars(query).all()
    for row, e in enumerate(events[:MAX_EVENTS], start=2):
        ts = e.ts if e.ts.tzinfo else e.ts.replace(tzinfo=timezone.utc)
        ws.cell(row, 1, ts.astimezone(timezone.utc).replace(tzinfo=None)).number_format = "yyyy-mm-dd hh:mm:ss"
        ws.cell(row, 2, EVENT_NAMES.get(e.type, e.type))
        ws.cell(row, 3, (e.anon_id or "")[:8])  # хватает различать устройства; целиком не нужен
        ws.cell(row, 4, e.user_id)
        ws.cell(row, 5, report.mask_ip(e.ip))
        ws.cell(row, 6, "да" if e.demo else "")
        ws.cell(row, 7, json.dumps(e.props, ensure_ascii=False) if e.props else "")
    if len(events) > MAX_EVENTS:
        ws.cell(MAX_EVENTS + 3, 1, f"Показаны первые {MAX_EVENTS:,} событий — сократите период".replace(",", " "))
    _widths(ws, [20, 18, 12, 13, 14, 7, 80])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{max(len(events[:MAX_EVENTS]) + 1, 2)}"


def _header(ws: Worksheet, row: int, names: list[str], first_col: int = 1) -> None:
    for col, name in enumerate(names, start=first_col):
        cell = ws.cell(row, col, name)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _widths(ws: Worksheet, widths: list[int]) -> None:
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
