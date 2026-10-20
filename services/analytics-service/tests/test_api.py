"""Тесты сервиса аналитики на SQLite, без Postgres и сервиса пользователей.

    PYTHONPATH=libs/wine-logging/src:services/analytics-service/src .venv/bin/python -m pytest services/analytics-service/tests -q
"""

from __future__ import annotations

import pytest
from fastapi.testclient import TestClient

from analytics_service.db import Database
from analytics_service.main import create_app
from analytics_service.report import mask_ip

WINES = [{"slug": "kuban-red", "title": "Пино Нуар", "manufacturer": "Гордиенко", "region": "Кубань"},
         {"slug": "crimea-white", "title": "Алиготе", "manufacturer": "Инкерман", "region": "Крым"}]


@pytest.fixture()
def client():
    state = {"admin": True}

    async def admin_check(request):
        return state["admin"]

    app = create_app(database=Database("sqlite://"), admin_check=admin_check, wine_source=lambda: WINES)
    with TestClient(app) as c:
        c.state = state
        yield c


def ev(client, type_, anon="anon-00000001", ip="95.165.12.40", **body):
    return client.post("/v1/events", json={"type": type_, "anon_id": anon, **body}, headers={"x-real-ip": ip})


def test_events_and_report(client):
    ev(client, "page_view")
    ev(client, "scan", props={"status": "match", "slug": "kuban-red", "title": "Пино Нуар", "region": "Кубань",
                              "manufacturer": "Гордиенко", "ms": 1500})
    ev(client, "scan", props={"status": "not_found", "ms": 2100})
    ev(client, "nudge_shown")
    ev(client, "nudge_click")
    ev(client, "register", user_id=7)
    ev(client, "share", user_id=7, props={"kind": "wine", "channel": "tg"})
    ev(client, "scan", anon="anon-00000002", ip="192.168.1.5", props={"status": "match", "slug": "kuban-red",
                                                                     "region": "Кубань"})
    r = client.get("/v1/admin/analytics?days=7").json()
    k = r["kpi"]
    assert (k["scans"], k["found"], k["not_found"], k["visitors"], k["registrations"], k["shares"]) == (3, 2, 1, 2, 1, 1)
    assert k["visitors_with_account"] == 1
    assert k["conversion"] == 0.5          # из двух сканировавших без аккаунта зарегистрировался один
    assert k["scan_p95_ms"] == 1500 or k["scan_p95_ms"] == 2100
    assert r["top_wines"] == [{"slug": "kuban-red", "title": "Пино Нуар", "count": 2}]
    assert r["top_regions"] == [{"name": "Кубань", "count": 2}]
    assert r["share_channels"] == [{"name": "tg", "count": 1}]
    assert [s["value"] for s in r["funnel"]] == [2, 1, 1, 1]
    assert {i["network"] for i in r["ips"]} == {"95.165.12.x", "192.168.1.x"}
    assert r["local_share"] == round(1 / 3, 3)
    assert len(r["daily"]) == 7 and sum(d["found"] + d["not_found"] for d in r["daily"]) == 3
    assert r["has_demo"] is False


def test_bad_events_rejected(client):
    assert ev(client, "hack").status_code == 422
    assert client.post("/v1/events", json={"type": "scan", "anon_id": "x"}).status_code == 422
    # сложные значения в props отбрасываются, строки обрезаются
    ev(client, "scan", props={"status": "match", "nested": {"a": 1}, "slug": "s" * 500})
    r = client.get("/v1/admin/analytics").json()
    assert r["kpi"]["scans"] == 1


def test_admin_only(client):
    client.state["admin"] = False
    assert client.get("/v1/admin/analytics").status_code == 403
    assert client.post("/v1/admin/analytics/demo").status_code == 403
    assert ev(client, "page_view").status_code == 204   # собирать события может любой


def test_demo_data(client):
    ev(client, "scan", props={"status": "match", "slug": "real"})
    created = client.post("/v1/admin/analytics/demo?days=14").json()["created"]
    assert created > 100
    r = client.get("/v1/admin/analytics?days=14").json()
    assert r["has_demo"] is True and r["kpi"]["scans"] > 100
    assert {w["slug"] for w in r["top_wines"]} <= {"kuban-red", "crimea-white", "real"}
    only_real = client.get("/v1/admin/analytics?days=14&demo=false").json()
    assert only_real["kpi"]["scans"] == 1 and only_real["has_demo"] is False
    # повторная генерация заменяет демо, удаление — убирает только демо
    client.post("/v1/admin/analytics/demo?days=14")
    assert client.delete("/v1/admin/analytics/demo").json()["deleted"] > 100
    assert client.get("/v1/admin/analytics?days=14").json()["kpi"]["scans"] == 1


def test_rate_limit(client, monkeypatch):
    monkeypatch.setattr("analytics_service.main.settings", type("S", (), {"max_events_per_minute": 3})())
    codes = [ev(client, "page_view", ip="10.0.0.9").status_code for _ in range(5)]
    assert codes == [204, 204, 204, 429, 429]


def test_mask_ip():
    assert mask_ip("95.165.12.40") == "95.165.12.x"
    assert mask_ip("2a00:1450:4010:c02::64") == "2a00:1450:4010::/48"
    assert mask_ip(None) == "неизвестно" and mask_ip("garbage") == "неизвестно"


def test_export_xlsx(client):
    from io import BytesIO

    from openpyxl import load_workbook

    ev(client, "scan", props={"status": "match", "slug": "kuban-red", "title": "Пино Нуар", "region": "Кубань"})
    ev(client, "scan", anon="anon-00000002", props={"status": "not_found"})
    ev(client, "share", props={"channel": "telegram", "kind": "wine"})
    res = client.get("/v1/admin/analytics/export.xlsx?days=7")
    assert res.status_code == 200
    assert res.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "attachment" in res.headers["content-disposition"]
    wb = load_workbook(BytesIO(res.content))
    assert wb.sheetnames == ["Сводка", "По дням", "Воронка", "Вина", "Регионы и производители",
                             "Поделиться", "Вовлечение", "Точность распознавания", "Оценки",
                             "Сети", "События"]
    kpi = {row[0]: row[1] for row in wb["Сводка"].iter_rows(min_row=6, values_only=True)}
    assert kpi["Сканов"] == 2 and kpi["Вино найдено"] == 1 and kpi["Поделились"] == 1
    assert wb["Вина"]["A3"].value == "Пино Нуар"
    assert wb["События"].max_row == 4  # заголовок + 3 события
    client.state["admin"] = False
    assert client.get("/v1/admin/analytics/export.xlsx").status_code == 403


def test_engagement_and_registration_sources(client):
    ev(client, "quiz", anon="anon-00000001", props={"correct": True})
    ev(client, "quiz", anon="anon-00000002", props={"correct": False})
    ev(client, "auth_prompt", anon="anon-00000001", props={"source": "quiz"})
    ev(client, "register", anon="anon-00000001")
    ev(client, "nudge_click", anon="anon-00000003", props={"place": "scan"})
    ev(client, "auth_prompt", anon="anon-00000003", props={"source": "diary"})  # последний призыв — дневник
    ev(client, "register", anon="anon-00000003")
    ev(client, "register", anon="anon-00000004")  # сам, без призыва
    ev(client, "auth_prompt", anon="anon-00000005", props={"source": "gift"})  # увидел, не зарегистрировался
    ev(client, "share", anon="anon-00000001", props={"channel": "gift", "kind": "gift"})
    ev(client, "gift_open", anon="anon-00000006")
    ev(client, "gift_save", anon="anon-00000006")
    ev(client, "diary", anon="anon-00000001", props={"kind": "tasted"})
    ev(client, "prize", anon="anon-00000001", props={"count": 2})
    g = client.get("/v1/admin/analytics?days=7").json()["engagement"]
    assert g["quiz"] == {"answers": 2, "players": 2, "correct_share": 0.5}
    assert g["gifts"] == {"created": 1, "opened": 1, "saved": 1} and g["prizes"] == 2
    src = {s["source"]: s for s in g["registration_sources"]}
    assert src["quiz"]["registrations"] == 1 and src["quiz"]["conversion"] == 1
    assert src["diary"]["registrations"] == 1 and src["scan"]["registrations"] == 0
    assert src["direct"]["registrations"] == 1 and src["gift"]["prompted"] == 1 and src["gift"]["conversion"] == 0


def test_demo_has_engagement(client):
    client.post("/v1/admin/analytics/demo?days=30")
    g = client.get("/v1/admin/analytics?days=30").json()["engagement"]
    assert g["quiz"]["answers"] > 0 and g["gifts"]["created"] > 0
    assert {s["source"] for s in g["registration_sources"]} >= {"scan", "quiz"}


def test_recognition_accuracy_and_ratings(client):
    for _ in range(3):
        ev(client, "scan", props={"status": "match", "slug": "a"})
    ev(client, "scan", props={"status": "not_found"})
    ev(client, "recognition", props={"answer": "confirmed", "status": "match"})
    ev(client, "recognition", props={"answer": "confirmed", "status": "match"})
    ev(client, "recognition", props={"answer": "picked_alternative", "status": "match"})
    ev(client, "recognition", props={"answer": "none_match", "status": "not_found"})
    ev(client, "diary", props={"kind": "tasted", "rating": 5})
    ev(client, "diary", props={"kind": "want", "rating": None})
    ev(client, "rate", props={"value": 1})
    ev(client, "rate", props={"value": -1})
    f = client.get("/v1/admin/analytics?days=7").json()["feedback"]
    assert f["answered"] == 3 and f["accuracy"] == 0.667
    week = f["weeks"][0]
    assert week["scans"] == 4 and week["match"]["confirmed"] == 2 and week["match"]["picked"] == 1
    assert week["not_found"] == {"scans": 1, "picked": 0, "missing": 1}
    assert f["stars"][0] == {"stars": 5, "count": 1} and f["likes"] == 1 and f["dislikes"] == 1
    from io import BytesIO

    from openpyxl import load_workbook

    book = load_workbook(BytesIO(client.get("/v1/admin/analytics/export.xlsx?days=7").content))
    ws = book["Точность распознавания"]
    assert ws["D5"].value == 2 and ws["G5"].value.startswith("=IF(C5>0")
    assert book["Оценки"]["B3"].value == 1
